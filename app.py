import os
import base64
import re
import io

from flask import Flask, request, jsonify, send_file, render_template
from mistralai import Mistral
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/process", methods=["POST"])
def process_pdf():
    if "file" not in request.files:
        return jsonify({"error": "Keine Datei übermittelt."}), 400

    file = request.files["file"]
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        return jsonify({"error": "Bitte nur PDF-Dateien hochladen."}), 400

    api_key = request.form.get("api_key") or os.getenv("MISTRAL_API_KEY", "")
    if not api_key:
        return jsonify({"error": "Kein API-Schlüssel angegeben."}), 400

    pdf_bytes = file.read()
    pdf_b64 = base64.standard_b64encode(pdf_bytes).decode("utf-8")

    try:
        client = Mistral(api_key=api_key)
        ocr_response = client.ocr.process(
            model="mistral-ocr-latest",
            document={
                "type": "document_url",
                "document_url": f"data:application/pdf;base64,{pdf_b64}",
            },
            include_image_base64=False,
        )
    except Exception as exc:
        return jsonify({"error": f"Mistral OCR Fehler: {exc}"}), 502

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    for page_idx, page in enumerate(ocr_response.pages):
        sheet_title = f"Seite {page_idx + 1}"
        ws = wb.create_sheet(title=sheet_title)
        _fill_sheet(ws, page.markdown)

    if not wb.sheetnames:
        ws = wb.create_sheet(title="Ergebnis")
        ws.append(["Kein Inhalt erkannt."])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = re.sub(r"\.pdf$", "", file.filename, flags=re.IGNORECASE) + ".xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_ALT_FILL = PatternFill("solid", fgColor="D9E1F2")
_THIN = Side(style="thin", color="AAAAAA")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_table(ws, start_row: int, nrows: int, ncols: int) -> None:
    for r in range(start_row, start_row + nrows):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r == start_row:
                cell.fill = _HEADER_FILL
                cell.font = _HEADER_FONT
            elif (r - start_row) % 2 == 0:
                cell.fill = _ALT_FILL

    for c in range(1, ncols + 1):
        col_letter = get_column_letter(c)
        ws.column_dimensions[col_letter].width = max(
            ws.column_dimensions[col_letter].width, 18
        )


def _parse_md_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if re.match(r"^\s*\|?[\s\-:|]+\|", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if any(cells):
            rows.append(cells)
    return rows


def _fill_sheet(ws, markdown: str) -> None:
    current_row = 1

    # Split into blocks separated by blank lines
    blocks: list[list[str]] = []
    current_block: list[str] = []
    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        if line == "":
            if current_block:
                blocks.append(current_block)
                current_block = []
        else:
            current_block.append(line)
    if current_block:
        blocks.append(current_block)

    for block in blocks:
        # Detect markdown table block (contains "|")
        if any("|" in l for l in block):
            table_rows = _parse_md_table(block)
            if table_rows:
                ncols = max(len(r) for r in table_rows)
                for row_data in table_rows:
                    # Pad short rows
                    padded = row_data + [""] * (ncols - len(row_data))
                    for col_idx, value in enumerate(padded, start=1):
                        ws.cell(row=current_row, column=col_idx, value=value)
                    current_row += 1
                _style_table(ws, current_row - len(table_rows), len(table_rows), ncols)
                current_row += 1  # blank row after table
                continue

        # Plain text block – write each line into column A
        for line in block:
            # Headings → bold
            heading_match = re.match(r"^(#{1,6})\s+(.*)", line)
            if heading_match:
                text = heading_match.group(2)
                cell = ws.cell(row=current_row, column=1, value=text)
                level = len(heading_match.group(1))
                cell.font = Font(bold=True, size=max(16 - level * 2, 10))
            else:
                ws.cell(row=current_row, column=1, value=line)
            current_row += 1

        current_row += 1  # blank row after text block

    # Auto-width column A
    max_len = 0
    for row in ws.iter_rows(min_col=1, max_col=1):
        for cell in row:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions["A"].width = min(max_len + 4, 80)


if __name__ == "__main__":
    app.run(debug=True, port=5000)
